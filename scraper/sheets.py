"""Reference spreadsheets (CSV) of everything the site uses, rebuilt after every scrape,
plus data/weights.csv: the hand-editable popularity weights.

data/sheets/restaurants.csv  every restaurant the planner can suggest
data/sheets/bars_venues.csv  bars, breweries, VFWs, venues, coffee shops...
data/sheets/activities.csv   bowling, pool, arcades, movies...
data/sheets/events.csv       every event we're tracking for the next two weeks
data/sheets/README.md        what each column means

weights.csv: 1 = normal, above 1 = suggest more often (2 = a strong boost),
below 1 = less often, 0 = never suggest. New places are added at 1 after each
scrape; numbers you've changed are never overwritten.
"""
from __future__ import annotations

import csv
import math
import re
from pathlib import Path

from .common import DATA, load_json

SHEETS = DATA / "sheets"
WEIGHTS = DATA / "weights.csv"
SITE = Path(__file__).resolve().parent.parent / "site"

AREAS = [("Downtown", 43.0745, -89.3895, 1.8), ("Willy St", 43.0835, -89.3655, 1.0), ("East side", 43.0960, -89.3420, 4.5),
         ("West side", 43.0660, -89.4550, 5.0), ("North side", 43.1350, -89.3600, 4.0), ("South side", 43.0300, -89.3950, 4.0),
         ("Middleton", 43.1000, -89.5050, 4.0), ("Sun Prairie", 43.1836, -89.2350, 5.0)]
CAPITOL = (43.0747, -89.3841)
GROUPS = {"trivia": "Games & trivia", "bingo": "Games & trivia", "karaoke": "Games & trivia",
          "music": "Music", "open_mic": "Music", "dance": "Music",
          "theater": "Theater", "film": "Theater", "comedy": "Comedy"}
CUISINE = {"american": "American", "latin": "Mexican & Latin", "italian": "Italian & Pizza", "asian": "Asian",
           "sushi": "Sushi & Japanese", "indian": "Indian", "med": "Mediterranean", "pub": "Bar & pub food"}


def miles(a, b):
    if a[0] is None or b[0] is None:
        return None
    r = math.radians
    h = math.sin(r(b[0] - a[0]) / 2) ** 2 + math.cos(r(a[0])) * math.cos(r(b[0])) * math.sin(r(b[1] - a[1]) / 2) ** 2
    return 2 * 3959 * math.asin(math.sqrt(h))


def area_of(lat, lon):
    if lat is None:
        return "No location"
    best, best_r = None, 1e9
    for name, la, lo, rkm in AREAS:
        if miles((lat, lon), (la, lo)) * 1.609 <= rkm < best_r:
            best, best_r = name, rkm
    if best:
        return best
    return "Out of town" if miles((lat, lon), CAPITOL) > 20 else "Outside the named areas"


def yn(v):
    return "" if v is None else "yes" if v else "no"


def hours_text(h):
    if not h:
        return ""
    days = "Mon Tue Wed Thu Fri Sat Sun".split()
    return "; ".join(",".join(days[d] for d in w[0]) + " " + w[1] + "-" + w[2] for w in h)


# ---------- weights ----------
def read_weights() -> dict[tuple[str, str], dict]:
    out = {}
    if WEIGHTS.exists():
        with WEIGHTS.open(newline="") as f:
            for row in csv.DictReader(f):
                if row.get("name"):
                    out[(row["kind"].strip(), row["name"].strip())] = row
    return out


def weight_of(w, kind, name):
    row = w.get((kind, name))
    try:
        return float(row["weight"]) if row and str(row["weight"]).strip() != "" else 1.0
    except ValueError:
        return 1.0


def update_weights(rest, venues, acts, events) -> dict:
    """Add every place and event venue we know about at weight 1; keep every number already set."""
    w = read_weights()
    seen = set()

    def add(kind, name, note=""):
        if not name or (kind, name) in seen:
            return
        seen.add((kind, name))
        w.setdefault((kind, name), {"kind": kind, "name": name, "weight": "1", "notes": note})
    for r in rest:
        add("restaurant", r["name"])
    for v in venues:
        add("bar" if v.get("kind") in ("bar", "brewery") else "venue", v["name"])
    for a in acts:
        add("activity", a["name"])
    for e in events:
        add("event venue", e.get("venue", ""))
    if not any(k == "event title" for k, _ in w):
        w[("event title", "Example: Comedy on State")] = {"kind": "event title", "name": "Example: Comedy on State", "weight": "1",
                                                         "notes": "any event whose title contains this text; delete or edit this example"}
    order = {"restaurant": 0, "bar": 1, "venue": 2, "activity": 3, "event venue": 4, "event title": 5}
    rows = sorted(w.values(), key=lambda r: (order.get(r["kind"], 9), r["name"].lower()))
    with WEIGHTS.open("w", newline="") as f:
        out = csv.DictWriter(f, fieldnames=["kind", "name", "weight", "notes"], extrasaction="ignore")
        out.writeheader()
        out.writerows(rows)
    return read_weights()


# ---------- sheets ----------
def write(name, cols, rows):
    with (SHEETS / name).open("w", newline="") as f:
        out = csv.writer(f)
        out.writerow(cols)
        out.writerows(rows)


def main():
    SHEETS.mkdir(parents=True, exist_ok=True)
    rest = load_json(DATA / "restaurants.json", {}).get("restaurants", [])
    venues = load_json(DATA / "venues.json", {}).get("venues", [])
    acts = load_json(DATA / "activities.json", {}).get("activities", [])
    events = load_json(DATA / "events.json", {}).get("events", [])
    vsrc = load_json(DATA / "venue_sources.json", {})
    picked = set(re.findall(r'\bn:"([^"]+)"', (SITE / "spots.js").read_text())) if (SITE / "spots.js").exists() else set()
    w = update_weights(rest, venues, acts, events)
    ev_count = {}
    for e in events:
        ev_count[e.get("venue_id") or e.get("venue")] = ev_count.get(e.get("venue_id") or e.get("venue"), 0) + 1
    ex = lambda s: (str(s or "").strip())

    write("restaurants.csv",
          ["name", "weight", "hand-picked", "cuisine group", "cuisine (as listed)", "suggested for dinner", "est. price per person",
           "fast/counter service", "serves alcohol", "alcohol info from", "area", "address", "menu link", "menu link from", "picture", "website", "hours",
           "latitude", "longitude", "place data from", "id"],
          [[r["name"], weight_of(w, "restaurant", r["name"]), yn(r["name"] in picked), CUISINE.get(r.get("group"), r.get("group") or ""),
            r.get("cuisine", ""), yn(r.get("dinner")), r.get("price"), yn(r.get("fast")), yn(r.get("alcohol")), r.get("alcohol_src", ""),
            area_of(r.get("lat"), r.get("lon")), r.get("address", ""),
            r.get("menu") or ("(broken: " + r["menu_broken"] + ")" if r.get("menu_broken") else "(none found: site links a Google menu search)"), r.get("menu_src", ""), r.get("image", ""), r.get("website", ""),
            hours_text(r.get("hours")), r.get("lat"), r.get("lon"), "OpenStreetMap", r["id"]]
           for r in sorted(rest, key=lambda r: r["name"].lower())])

    def feeds(v):
        f = v.get("feeds") or (vsrc.get(v["id"]) or {}).get("feeds")
        return "; ".join((x.get("type", "") + " " + x.get("url", "")).strip() if isinstance(x, dict) else str(x) for x in (f or []))
    write("bars_venues.csv",
          ["name", "kind", "weight", "events in next 2 weeks", "serves alcohol", "alcohol info from", "area", "address", "city",
           "website", "picture", "facebook", "instagram", "event feeds we read", "latitude", "longitude", "place data from", "id"],
          [[v["name"], v.get("kind", ""), weight_of(w, "bar" if v.get("kind") in ("bar", "brewery") else "venue", v["name"]),
            ev_count.get(v["id"], 0), yn(v.get("alcohol")), v.get("alcohol_src", ""), area_of(v.get("lat"), v.get("lon")),
            v.get("address", ""), v.get("city", ""), v.get("website", ""), v.get("image", ""), v.get("facebook", ""), v.get("instagram", ""), feeds(v),
            v.get("lat"), v.get("lon"), "OpenStreetMap" if str(v["id"]).startswith("osm:") else "added by hand / event listing", v["id"]]
           for v in sorted(venues, key=lambda v: (v.get("kind", ""), v["name"].lower()))])

    write("activities.csv",
          ["name", "type", "weight", "est. price per person", "typical minutes", "serves alcohol", "alcohol info from", "area",
           "address", "website", "picture", "hours", "latitude", "longitude", "place data from", "id"],
          [[a["name"], a.get("label", a.get("type", "")), weight_of(w, "activity", a["name"]), a.get("price"), a.get("minutes"),
            yn(a.get("alcohol")), a.get("alcohol_src", ""), area_of(a.get("lat"), a.get("lon")), a.get("address", ""),
            a.get("website", ""), a.get("image", ""), hours_text(a.get("hours")), a.get("lat"), a.get("lon"), "OpenStreetMap", a["id"]]
           for a in sorted(acts, key=lambda a: a["name"].lower())])

    def ev_weight(e):
        x = weight_of(w, "event venue", e.get("venue", ""))
        for (k, name), _ in w.items():
            if k == "event title" and not name.startswith("Example:") and name.lower() in e.get("title", "").lower():
                x *= weight_of(w, k, name)
        return round(x, 2)
    write("events.csv",
          ["date", "start", "end", "title", "list category", "detailed category", "venue", "venue kind", "area", "address",
           "price", "price is", "weight", "link", "link check", "picture", "found by", "all sources", "listed by", "description", "latitude", "longitude", "id"],
          [[ex(e.get("start"))[:10], "all day" if e.get("all_day") else ex(e.get("start"))[11:16], ex(e.get("end"))[11:16],
            e.get("title", ""), GROUPS.get(e.get("category"), "Other (All only)"), e.get("category", ""), e.get("venue", ""),
            e.get("venue_kind", ""), area_of(e.get("lat"), e.get("lon")), e.get("address", ""),
            "" if e.get("price") is None else e["price"], "from the listing" if e.get("price") is not None else "estimated on site",
            ev_weight(e), e.get("url", "") if "isthmus.com" not in ex(e.get("url")) else "",
            {True: "works", False: "broken (site uses a search)", None: ""}.get(e.get("link_ok"), ""), e.get("image", ""), e.get("source", ""),
            "; ".join(e.get("sources") or []), e.get("via", ""), ex(e.get("description"))[:500], e.get("lat"), e.get("lon"), e.get("id", "")]
           for e in sorted(events, key=lambda e: (ex(e.get("start")), e.get("title", "")))])

    # music venues: how many events we find at each, and a place to add a calendar link
    cal_path = DATA / "venue_calendars.csv"
    cals = {}
    if cal_path.exists():
        with cal_path.open(newline="") as fh:
            cals = {r["name"]: r for r in csv.DictReader(fh) if r.get("name")}
    music = {}
    for e in events:
        if e.get("category") in ("music", "open_mic", "dance", "karaoke"):
            music[e.get("venue_id")] = music.get(e.get("venue_id"), 0) + 1
    rows = []
    for v in venues:
        is_music = (v.get("kind") in ("venue", "bar", "brewery") and (music.get(v["id"]) or re.search(
            r"music|saloon|cabaret|tavern|pub\b|lounge|sylvee|majestic|bur oak|cardinal|rigby|liquid|harmony|gamma ray|up north|knuckle|red rooster|wisco|tip top|mickey|crystal corner|high noon|cafe coda|robinia|bos meadery|winnebago|come back in|brink|mr\.? robert", v["name"], re.I)))
        if not is_music:
            continue
        src = vsrc.get(v["id"]) or {}
        n_all = ev_count.get(v["id"], 0)
        status = ("calendar link added" if v["name"] in cals and cals[v["name"]].get("calendar_url") else
                  "no website" if not v.get("website") else "site unreachable" if src.get("error") else
                  "reading OK" if n_all else "website has no readable calendar")
        rows.append([v["name"], v.get("kind", ""), music.get(v["id"], 0), n_all, status, v.get("website", ""),
                     (cals.get(v["name"]) or {}).get("calendar_url", ""), (cals.get(v["name"]) or {}).get("notes", "")])
        cals.setdefault(v["name"], {"name": v["name"], "calendar_url": "", "notes": ""})
    rows.sort(key=lambda r: (r[3] > 0, r[0].lower()))
    write("music_venues.csv", ["name", "kind", "music events (2 weeks)", "all events (2 weeks)", "status", "website",
                               "calendar_url", "notes"], rows)
    with cal_path.open("w", newline="") as fh:  # keeps every link and note already filled in
        cw = csv.DictWriter(fh, fieldnames=["name", "calendar_url", "notes"], extrasaction="ignore")
        cw.writeheader()
        cw.writerows(sorted(cals.values(), key=lambda r: (not r.get("calendar_url"), r["name"].lower())))

    (SHEETS / "README.md").write_text(README)
    try:
        write_xlsx()
    except Exception as e:  # the CSVs are still there
        print(f"Workbook skipped ({type(e).__name__}: {e})")
    print(f"Sheets: {len(rest)} restaurants, {len(venues)} bars & venues, {len(acts)} activities, {len(events)} events; "
          f"{len(w)} weight rows")


def write_xlsx():
    """One workbook with every sheet, for opening in Excel or Google Sheets: data/sheets/MadisonPulse.xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head = PatternFill("solid", fgColor="14202B")
    fill_in = PatternFill("solid", fgColor="FFF2A8")
    link = Font(name="Arial", size=10, color="1F5A7A", underline="single")
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = [
        ("MadisonPulse reference workbook", True),
        ("Rebuilt automatically after every scrape (about 6 am and 6 pm). Latest copy: data/sheets/MadisonPulse.xlsx in the GitHub repo.", False),
        ("", False),
        ("Dead links: fill in the yellow New link column (and Notes if you like), then either edit data/link_fixes.csv on GitHub", False),
        ("  or send this workbook back to Claude. The next scrape uses your new link everywhere that dead link appeared.", False),
        ("  Until then the site shows \"Link down\" for it. A dead website can also mean the place has closed.", False),
        ("Weights: change data/weights.csv on GitHub (1 = normal, 1.5 = more often, 2 = a lot more, 0.5 = less, 0 = never).", False),
        ("  The weight shown on each tab here is read from that file.", False),
        ("", False),
        ("Music venues: places that host music. Ones with few or no events usually post them somewhere we can't read.", False),
        ("  Fill in the yellow calendar_url with the page that lists their shows (or send the workbook back); the next scrape reads it.", False),
        ("", False),
        ("Tabs: Dead links, Music venues, Restaurants, Bars & venues, Activities, Events, Weights.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = readme.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=14 if b else 10, bold=b)
    readme.column_dimensions["A"].width = 120

    def add(title, path, fill_cols=(), example=None):
        if not path.exists():
            return
        with path.open(newline="") as f:
            rows = list(csv.reader(f))
        if not rows:
            return
        ws = wb.create_sheet(title)
        cols = rows[0]
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = bold, head
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows[1:]:
            ws.append([float(v) if v and re.fullmatch(r"-?\d+(\.\d+)?", v) and len(v) < 12 else v for v in r])
        for col_i, name in enumerate(cols, 1):
            width = max([len(str(name))] + [len(str(r[col_i - 1])) for r in rows[1:200] if len(r) >= col_i])
            ws.column_dimensions[get_column_letter(col_i)].width = min(max(10, width + 2), 48)
            fill = name in fill_cols
            for cell in ws.iter_cols(min_col=col_i, max_col=col_i, min_row=2, max_row=ws.max_row).__next__():
                v = cell.value
                if isinstance(v, str) and v.startswith(("http://", "https://")):
                    cell.hyperlink, cell.font = v, link
                else:
                    cell.font = font
                if fill:
                    cell.fill = fill_in
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    add("Dead links", DATA / "link_fixes.csv", fill_cols=("new_url", "notes"))
    add("Music venues", SHEETS / "music_venues.csv", fill_cols=("calendar_url", "notes"))
    add("Restaurants", SHEETS / "restaurants.csv")
    add("Bars & venues", SHEETS / "bars_venues.csv")
    add("Activities", SHEETS / "activities.csv")
    add("Events", SHEETS / "events.csv")
    add("Weights", WEIGHTS)
    if "Dead links" in wb.sheetnames:
        ws = wb["Dead links"]
        for c in ws[1]:
            c.value = {"dead_url": "Dead link", "new_url": "New link (fill in)", "what": "What it is", "name": "Place / event",
                       "why": "Why it's dead", "first_seen": "First seen", "notes": "Notes (fill in)"}.get(c.value, c.value)
    wb.save(SHEETS / "MadisonPulse.xlsx")


README = """# MadisonPulse reference sheets

Rebuilt automatically after every scrape (about 6 am and 6 pm). Click a file to see it as a
searchable table on GitHub, or use the download button to open it in Excel or Google Sheets.

| File | What's in it |
| --- | --- |
| restaurants.csv | Every restaurant the planner can suggest: cuisine, price estimate, alcohol, menu link, weight |
| bars_venues.csv | Bars, breweries, VFWs, lodges, venues, coffee shops: event feeds we read, events found, alcohol |
| activities.csv | Bowling, pool, arcades, movies and similar |
| events.csv | Every event in the next two weeks: category, price (listed or estimated), link, which sources found it |

**Weights** live in `data/weights.csv` (one level up). Change the `weight` column to steer
suggestions: 1 = normal, 1.5 = more often, 2 = a lot more, 0.5 = less often, 0 = never.
`event title` rows match any event whose title contains that text. New places are added at 1
after each scrape; numbers you change are never overwritten. The site picks up an edited
weights file within a couple of minutes.

**Where things come from**
- Places: OpenStreetMap (national chains left out), refreshed weekly.
- Alcohol: an OpenStreetMap tag, a name/type rule (e.g. "tavern", "brewery", juice bars), or an
  AI read of the place's own website.
- Menu links: OpenStreetMap's menu tag, or the "Menu" link on the restaurant's homepage
  (checked monthly). Without one, the site links a Google search for the menu.
- Pictures: the image each place or event page publishes for link previews (what shows when
  it's shared on Facebook or in a text), from their own site, the event feed or Ticketmaster.
  Shown from their site, never copied; nothing from Isthmus.
- Events: venue websites and calendars, the Isthmus city calendar (facts only), Ticketmaster,
  and AI reading of venue event pages. "all sources" lists every source that had the event.
"""

if __name__ == "__main__":
    main()
